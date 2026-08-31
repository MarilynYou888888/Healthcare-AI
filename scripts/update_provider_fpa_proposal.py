from pathlib import Path

from docx import Document


SOURCE = Path("VCP_AI_Project_Proposal_Filled_Healthcare_Market_Intelligence.docx")
OUTPUT = Path("VCP_AI_Project_Proposal_Provider_FPA_Variance_Copilot.docx")


def replace_preserving_first_run(paragraph, text):
    if not paragraph.runs:
        paragraph.add_run(text)
        return
    paragraph.runs[0].text = text
    for run in paragraph.runs[1:]:
        run._element.getparent().remove(run._element)


def set_label(paragraph, label, body):
    for run in list(paragraph.runs):
        run._element.getparent().remove(run._element)
    label_run = paragraph.add_run(label)
    label_run.bold = True
    paragraph.add_run(body)


def set_bullet(paragraph, label, body):
    for run in list(paragraph.runs):
        run._element.getparent().remove(run._element)
    label_run = paragraph.add_run("• " + label)
    label_run.bold = True
    paragraph.add_run(body)


document = Document(SOURCE)
paragraphs = document.paragraphs

replace_preserving_first_run(paragraphs[0], "AI Project Proposal")
replace_preserving_first_run(paragraphs[1], "Provider FP&A Variance Copilot")
replace_preserving_first_run(
    paragraphs[2],
    "An evidence-backed variance explanation and forecast-assumption review workflow for U.S. healthcare providers",
)

set_label(
    paragraphs[4],
    "Product Name / One-Line Pitch: ",
    "Provider FP&A Variance Copilot helps financial analysts at U.S. hospital systems and multi-site specialty providers identify material actual-versus-forecast or actual-versus-budget variances, connect them to supported operating drivers, and prepare analyst-reviewable forecast-assumption change proposals.",
)

set_label(
    paragraphs[6],
    "The Problem: ",
    "Provider FP&A teams spend substantial recurring time building facility forecasts, investigating revenue and expense variances, reconciling leadership assumptions, and gathering operating context from finance, operations, accounting, revenue cycle, and BI. A numerical variance is visible before its explanation is: analysts must determine whether the driver is temporary, structural, an accounting-timing issue, or still unresolved before deciding whether a forecast assumption deserves review.",
)
set_label(
    paragraphs[7],
    "Target User: ",
    "A Provider FP&A Analyst responsible for recurring forecasts, budgets, variance analysis, and facility, market, or service-line performance at a U.S. hospital system or multi-site specialty provider.",
)
set_label(
    paragraphs[8],
    "Intended Impact / Customer-Discovery Basis: ",
    "Three recent interviews with finance and operations professionals at Ardent Health, Upperline Health, and Monogram Health surfaced a concrete recurring need: explain performance at the facility or market level, distinguish temporary volume issues from operational problems, and prepare decisions for analyst and operational review. The MVP aims to shorten that investigation while improving evidence traceability and consistency; it does not replace financial judgment.",
)

set_label(
    paragraphs[10],
    "Primary Organization: ",
    "A U.S. healthcare provider organization, initially a hospital system or multi-site specialty provider.",
)
set_label(
    paragraphs[11],
    "User Profile: ",
    "The user is an FP&A or Financial Analyst who works across Excel or a planning system, ERP or general-ledger exports, BI dashboards, operational reports, and recurring forecast calendars. The analyst partners with operations, accounting, revenue cycle, and leadership and remains accountable for the final interpretation and forecast decision.",
)
set_label(
    paragraphs[12],
    "Usage Scenario: ",
    "During a recurring forecast or variance review, the analyst selects an assigned hospital, clinic, market, or service line and a reporting period. The workflow identifies a material patient-volume variance, checks available operating context such as provider PTO, temporary clinic closure, weather disruption, capacity change, new-provider ramp, reimbursement or mix change, and accounting timing, then shows which candidate explanations are supported. It proposes whether the relevant forecast assumption should remain unchanged, be scenario-tested, or be reviewed further.",
)

set_label(paragraphs[14], "Core Features (MVP - Must-Have): ", "")
set_bullet(
    paragraphs[15],
    "Variance detection and prioritization: ",
    "compare actual results with forecast or budget by metric or account, reporting period, and facility, market, or service line; surface only material variances using analyst-controlled thresholds.",
)
set_bullet(
    paragraphs[16],
    "Operating-context assembly: ",
    "bring together available operational metrics, schedules, capacity and closure information, reimbursement or mix data, accounting-close context, prior assumptions, and analyst-provided notes without inventing missing evidence.",
)
set_bullet(
    paragraphs[17],
    "Evidence-backed driver assessment: ",
    "rank candidate Operating Drivers, identify the Supporting Evidence for and against each one, classify timing as temporary, structural, or unresolved, and expose unanswered questions.",
)
set_bullet(
    paragraphs[18],
    "Assumption-Change Proposal: ",
    "recommend retain, scenario-test, or review, and include the metric or account, variance, Performance Unit, candidate Operating Driver, Supporting Evidence, financial impact mechanism, Timing Classification, confidence, unresolved questions, and Analyst Approval status.",
)

set_label(paragraphs[19], "Phase 2 - Retained Market Intelligence Module: ", "")
set_bullet(
    paragraphs[20],
    "External-event monitoring: ",
    "preserve the Healthcare Market Intelligence Copilot concept as a later module that monitors approved healthcare news, filings, regulatory developments, research, and user-provided sources.",
)
set_bullet(
    paragraphs[21],
    "External-to-financial-driver mapping: ",
    "connect material external developments to provider-relevant drivers such as patient volume, reimbursement, payer mix, labor, supply expense, and capital plans, with visible citations and uncertainty.",
)
set_bullet(
    paragraphs[22],
    "Later workflow extensions: ",
    "consider alerts, dashboards, watchlists, ROI-analysis support, and invoice or expense workflows only after the Provider FP&A MVP demonstrates repeatable value.",
)

set_label(
    paragraphs[24],
    "Product Format: ",
    "An analyst-facing AI workflow used during recurring variance and forecast review. The initial version may operate through a structured agent or lightweight web interface and should accept governed file exports before attempting deeper enterprise integrations.",
)
set_label(
    paragraphs[25],
    "Proposed Architecture: ",
    "Use deterministic calculations for actual-versus-forecast and actual-versus-budget variances; a governed data layer for finance, operating context, assumptions, source lineage, and review state; and an LLM analysis layer for candidate-driver assessment and structured explanation. Integration design should favor read-only access, least privilege, auditable transformations, and deployment within the provider's approved environment. Specific model, database, and cloud choices remain provisional until data-access and security constraints are validated.",
)

set_label(
    paragraphs[27],
    "Data Needs: ",
    "Actual, forecast, and budget values by account or metric, period, and Performance Unit; operational KPIs and capacity context; provider schedules and PTO when available; clinic closure and disruption records; reimbursement and payer-mix indicators; accounting-close and timing notes; current Forecast Assumptions; materiality thresholds; and Analyst Approval history.",
)
set_label(
    paragraphs[28],
    "Initial Data Sources: ",
    "Start with user-provided, governed exports from the ERP or general ledger, planning and budgeting tools, BI dashboards, revenue-cycle reports, scheduling or workforce systems, and operational notes. Each value or explanation must retain its source, reporting period, Performance Unit, extraction time, and transformation history. External sources are optional context for the MVP rather than its primary input.",
)
set_label(
    paragraphs[29],
    "Access, Privacy, and Cost: ",
    "Use the minimum necessary data, avoid protected health information where the workflow does not require it, and separate source access from model access. Begin with read-only file-based inputs to validate value and governance. Any later system integration will require organization-specific authentication, authorization, retention, audit, security, and vendor-review decisions.",
)

set_label(
    paragraphs[31],
    "Role of AI/LLM: ",
    "AI organizes the investigation: it links a Variance to candidate Operating Drivers, retrieves Supporting Evidence, explains the financial impact mechanism, classifies timing, reports confidence and unresolved questions, and drafts an Assumption-Change Proposal. Deterministic logic performs financial calculations; the analyst evaluates the explanation and makes the decision.",
)
set_label(
    paragraphs[32],
    "Human-Control Boundary and Implementation Depth: ",
    "The system never directly updates an ERP, general ledger, forecast, or budget. It produces structured proposals with source lineage and an explicit Analyst Approval status. The MVP should use schemas and rules that require every proposal field, reject unsupported conclusions, distinguish missing evidence from negative evidence, and maintain an evaluation set drawn from historical variance cases. Fine-tuning is not required for the first version.",
)

set_label(paragraphs[34], "Key Performance Indicators (KPIs): ", "")
set_bullet(
    paragraphs[35],
    "Analyst time saved: ",
    "reduce recurring variance-investigation and forecast-review preparation time by at least 40% on pilot cases without increasing review burden.",
)
set_bullet(
    paragraphs[36],
    "Calculation and traceability: ",
    "100% of reported variance values reconcile to the governed input data within the agreed tolerance, and 100% of proposed Operating Drivers link to visible Supporting Evidence or are explicitly labeled unresolved.",
)
set_bullet(
    paragraphs[37],
    "Explanation quality: ",
    "on a blinded historical test set, analysts rate at least 80% of prioritized candidate drivers as relevant and identify no unsupported high-confidence conclusion in an approved output.",
)
set_bullet(
    paragraphs[38],
    "Decision safety and adoption: ",
    "zero automatic writes to ERP, general-ledger, forecast, or budget systems; every proposal records Analyst Approval; and pilot analysts use the workflow in at least four consecutive recurring review cycles.",
)

replace_preserving_first_run(
    paragraphs[40],
    "This project helps Provider FP&A Analysts at U.S. hospital systems and multi-site specialty providers explain material financial variances, connect them to evidence-backed Operating Drivers, and prepare analyst-reviewable Forecast Assumption changes without automatically modifying the financial model; external healthcare market intelligence remains a possible Phase 2 module.",
)

document.save(OUTPUT)
print(OUTPUT.resolve())
